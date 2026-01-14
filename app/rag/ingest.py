import os
from typing import List
from uuid import uuid4

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from pypdf import PdfReader
from pptx import Presentation

from app.config import settings
from app.rag.vectorstore import get_vectorstore


def _load_pdf(path: str) -> List[Document]:
    reader = PdfReader(path)
    docs: List[Document] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": os.path.basename(path), "page": i + 1},
                )
            )
    return docs


def _load_pptx(path: str) -> List[Document]:
    prs = Presentation(path)
    docs: List[Document] = []
    for idx, slide in enumerate(prs.slides):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                t = shape.text.strip()
                if t:
                    parts.append(t)
        text = "\n".join(parts).strip()
        if text:
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": os.path.basename(path), "page": idx + 1},
                )
            )
    return docs


def load_documents(path: str) -> List[Document]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".pptx":
        return _load_pptx(path)
    raise ValueError(f"Unsupported file type: {ext} (only PDF/PPTX supported)")


def ingest_file(filepath: str) -> int:
    raw_docs = load_documents(filepath)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900,
        chunk_overlap=150,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(raw_docs)

    for c in chunks:
        c.metadata = dict(c.metadata or {})
        c.metadata["chunk_id"] = str(uuid4())

    vs = get_vectorstore()
    if chunks:
        vs.add_documents(chunks)
        vs.persist()

    return len(chunks)
